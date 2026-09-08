"""文档加载器 - 支持txt/doc/pdf/markdown"""
import os
import re
from pathlib import Path
from typing import List, Dict


def load_documents(file_path: str, category: str=None) -> List[Dict]:
    """
    加载并解析文档，返回结构化文档块列表
    每个文档块: {"content": str, "metadata": dict}

    Args:
        file_path: 文档路径
        category: 文档分类（如"临床指南"、"药物信息"等，从目录名推断）
    """
    ext = Path(file_path).suffix.lower()
    file_name = Path(file_path).stem  # 不含扩展名的文件名

    # 从路径推断分类（如果没传）
    if not category:
        parts = Path(file_path).parts
        for part in parts:
            if part in ("临床指南", "疾病症状与诊疗", "药物信息"):
                category = part
                break

    base_meta = {
        "file_name": Path(file_path).name,
        "title": file_name,  # 默认标题为文件名
        "category": category or "其他",
        "source": file_path,
    }

    if ext in (".txt",):
        return _load_txt(file_path,base_meta)
    elif ext in (".md", ".markdown"):
        return _load_markdown(file_path, base_meta)
    elif ext == ".pdf":
        return _load_pdf(file_path,base_meta)
    elif ext in (".doc", ".docx"):
        return _load_docx(file_path,base_meta)
    else:
        raise ValueError(f"不支持的文件类型: {ext}")


# ============================================================
# 文本清洗工具
# ============================================================

def clean_text(text: str) -> str:
    """文本清洗流水线"""
    if not text:
        return ""

    # 1. 统一换行符
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # 2. 去除行首尾空白
    lines = [line.strip() for line in text.split("\n")]

    # 3. 合并多余空行（最多保留1个空行）
    cleaned_lines = []
    prev_empty = False
    for line in lines:
        if not line:
            if not prev_empty:
                cleaned_lines.append("")
            prev_empty = True
        else:
            # 去除行内多余空格（中文之间的空格保留1个）
            line = re.sub(r"[ \t]+", " ", line)
            cleaned_lines.append(line)
            prev_empty = False

    text = "\n".join(cleaned_lines)

    # 4. 去除不可见字符（保留常用空白和标点）
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # 5. 去除常见页眉页脚噪音（简单启发式）
    text = _remove_header_footer_noise(text)

    return text.strip()

def _remove_header_footer_noise(text: str) -> str:
    """去除常见页眉页脚噪音（简单规则，后续可优化）"""
    lines = text.split("\n")
    filtered = []
    for line in lines:
        stripped = line.strip()
        # 跳过纯数字行（页码）
        if re.fullmatch(r"\d+", stripped):
            continue
        # 跳过"第X页 共Y页"格式
        if re.search(r"第\s*\d+\s*页", stripped) and re.search(r"共\s*\d+\s*页", stripped):
            continue
        # 跳过只有"---""或"==="等分隔线
        if re.fullmatch(r"[-=—_…·.。\s]+", stripped) and len(stripped) > 3:
            continue
        filtered.append(line)
    return "\n".join(filtered)



def _load_txt(file_path: str, base_meta: dict) -> List[Dict]:
    """加载纯文本文件"""
    text = _read_file_with_encoding(file_path)
    text = clean_text(text)

    # 尝试从首行提取标题
    lines = text.split("\n")
    title = base_meta["title"]
    if lines and lines[0].strip():
        first_line = lines[0].strip()
        # 如果首行较短且不是正文开头，可能是标题
        if len(first_line) < 50 and not first_line.endswith(("。", "！", "？", ".")):
            title = first_line
            text = "\n".join(lines[1:]).strip()

    meta = {**base_meta, "title": title}
    return [{"content": text, "metadata": meta}]

def _read_file_with_encoding(file_path: str) -> str:
    """尝试多种编码读取文本文件"""
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码文件: {file_path}")


def _load_markdown(file_path: str, base_meta: dict) -> List[Dict]:
    """
    结构化加载 Markdown 文档
    按标题层级拆分，每个块带上完整的标题路径（breadcrumb）
    例如：高血压 -> 症状 -> 常见症状
    """
    from markdown_it import MarkdownIt

    raw_text = _read_file_with_encoding(file_path)
    raw_text = clean_text(raw_text)

    md = MarkdownIt("commonmark")
    tokens = md.parse(raw_text)

    # 提取文档标题（第一个一级标题，或文件名）
    doc_title = base_meta["title"]
    for idx, tok in enumerate(tokens):
        if tok.type == "heading_open" and tok.tag == "h1":
            if idx + 1 < len(tokens) and tokens[idx + 1].type == "inline":
                doc_title = tokens[idx + 1].content.strip()
            break

    # 按标题层级拆分文档
    sections = _split_markdown_by_headings(raw_text, doc_title)

    # 组装结果
    results = []
    for section in sections:
        heading_path = section["headings"]  # 标题路径列表
        content = section["content"].strip()
        if not content:
            continue

        # 构造完整标题（面包屑格式）
        full_title = " - ".join(heading_path)

        # 在内容前加上标题，增强语义
        content_with_title = f"# {full_title}\n\n{content}"

        meta = {
            **base_meta,
            "title": doc_title,
            "section_title": full_title,  # 章节完整标题
            "heading_level": len(heading_path),  # 标题层级深度
            "heading_path": " > ".join(heading_path),  # 面包屑字符串
        }
        results.append({"content": content_with_title, "metadata": meta})

    # 如果没解析出任何章节（说明md没有标题结构），返回整块
    if not results:
        meta = {**base_meta, "title": doc_title}
        results = [{"content": raw_text, "metadata": meta}]

    return results


def _split_markdown_by_headings(text: str, doc_title: str) -> List[Dict]:
    """
    按 Markdown 标题层级拆分为多个 section
    每个 section 包含: headings(标题路径列表), content(该节内容)
    """
    lines = text.split("\n")
    sections = []
    heading_stack = []  # 当前标题栈：[(level, text)]
    current_content = []
    current_headings = [doc_title]  # 初始以文档标题为根

    heading_pattern = re.compile(r'^(#{1,6})\s+(.+)$')

    def _save_section():
        content = "\n".join(current_content).strip()
        if content or heading_stack:
            sections.append({
                "headings": current_headings.copy(),
                "content": content,
            })

    for line in lines:
        match = heading_pattern.match(line.strip())
        if match:
            # 遇到新标题，先保存当前节
            if current_content or heading_stack:
                _save_section()

            level = len(match.group(1))
            title_text = match.group(2).strip()

            # 弹出同级或更深的标题
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()

            heading_stack.append((level, title_text))
            current_headings = [doc_title] + [h[1] for h in heading_stack]
            current_content = []
        else:
            current_content.append(line)

    # 保存最后一节
    if current_content or heading_stack:
        _save_section()

    return sections


# ============================================================
# PDF 加载（核心优化点：改用 pdfplumber）
# ============================================================

def _load_pdf(file_path: str, base_meta: dict) -> List[Dict]:
    """
    使用 pdfplumber 加载 PDF
    - 更好的文本排版还原
    - 表格提取为 Markdown 表格格式
    - 每页独立返回，保留页码元数据
    """
    import pdfplumber

    results = []
    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)

        for page_num, page in enumerate(pdf.pages, start=1):
            page_text = ""

            # 1. 提取页面文本
            text = page.extract_text() or ""

            # 2. 提取表格，转为 Markdown 表格
            tables = page.extract_tables()
            if tables:
                table_texts = []
                for table in tables:
                    md_table = _table_to_markdown(table)
                    if md_table:
                        table_texts.append(md_table)

                # 如果有表格，把表格内容附在页面文本后面
                if table_texts:
                    text = text + "\n\n" + "\n\n".join(table_texts)

            text = clean_text(text)
            if not text:
                continue

            # 尝试从第一页提取文档标题
            title = base_meta["title"]
            if page_num == 1:
                first_lines = [l.strip() for l in text.split("\n") if l.strip()]
                if first_lines and len(first_lines[0]) < 100:
                    title = first_lines[0]

            meta = {
                **base_meta,
                "title": title,
                "page": page_num,
                "total_pages": total_pages,
            }
            results.append({"content": text, "metadata": meta})

    # 如果整本书没提取出内容（可能是扫描版），返回提示
    if not results:
        meta = {**base_meta}
        results = [{"content": "（该PDF可能为扫描版，无法提取文本内容）", "metadata": meta}]

    return results


def _table_to_markdown(table: List[List]) -> str:
    """将二维表格列表转为 Markdown 表格字符串"""
    if not table or not table[0]:
        return ""

    # 清理单元格内容
    cleaned = []
    for row in table:
        cleaned_row = [
            str(cell).strip().replace("\n", "<br>") if cell else ""
            for cell in row
        ]
        # 跳过全空行
        if any(cell.strip() for cell in cleaned_row):
            cleaned.append(cleaned_row)

    if not cleaned:
        return ""

    # 第一行作为表头
    header = cleaned[0]
    body = cleaned[1:] if len(cleaned) > 1 else []

    # 构造 Markdown 表格
    md_lines = []
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in body:
        # 补齐列数
        if len(row) < len(header):
            row = row + [""] * (len(header) - len(row))
        md_lines.append("| " + " | ".join(row[:len(header)]) + " |")

    return "\n".join(md_lines)


# ============================================================
# Word 文档加载
# ============================================================

def _load_docx(file_path: str, base_meta: dict) -> List[Dict]:
    """加载 Word 文档（docx 用 python-docx，doc 暂不支持）"""
    ext = Path(file_path).suffix.lower()

    if ext == ".docx":
        from docx import Document
        doc = Document(file_path)

        # 提取段落
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # 提取表格（转为 Markdown 表格）
        table_texts = []
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row.cells])
            md_table = _table_to_markdown(table_data)
            if md_table:
                table_texts.append(md_table)

        full_text = "\n".join(paragraphs)
        if table_texts:
            full_text += "\n\n" + "\n\n".join(table_texts)

        full_text = clean_text(full_text)

        # 提取标题（第一段）
        title = base_meta["title"]
        if paragraphs and len(paragraphs[0]) < 100:
            title = paragraphs[0].strip()

        meta = {**base_meta, "title": title}
        return [{"content": full_text, "metadata": meta}]
    else:
        # .doc 格式暂不支持完整解析
        with open(file_path, "rb") as f:
            raw = f.read()
        try:
            text = raw.decode("utf-8", errors="ignore")
            text = "".join(c for c in text if c.isprintable() or c in "\n\r\t")
            text = clean_text(text)
            meta = {**base_meta}
            return [{"content": text, "metadata": meta}]
        except Exception:
            raise ValueError(f"无法解析.doc文件: {file_path}")


# ============================================================
# 兼容旧接口：返回纯文本字符串
# ============================================================

def load_document(file_path: str) -> str:
    """
    兼容旧版接口：加载文档并返回纯文本字符串
    新代码请使用 load_documents() 获取结构化结果
    """
    docs = load_documents(file_path)
    return "\n\n".join(doc["content"] for doc in docs)